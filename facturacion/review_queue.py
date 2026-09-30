"""Cola de revisión manual: casos donde el script no tiene suficiente certeza
para registrar automáticamente (solo PDF/asunto sin XML, nota crédito sin
factura identificada, posible duplicado no confirmado, etc.).

Vive en un Excel aparte (no en el Consolidado 2026.xlsx oficial) para que
nada se registre en el archivo de producción sin que alguien lo confirme,
pero tampoco se pierda de vista.
"""

import os

from openpyxl import Workbook, load_workbook

from . import excel_store

ENCABEZADOS_PENDIENTES = [
    "Fecha detección",
    "Hora recepción correo",
    "Remitente",
    "Asunto",
    "Categoría",
    "Motivo de la duda",
    "Factura (tentativa)",
    "Nit (tentativo)",
    "Proveedor (tentativo)",
    "Valor sin iva (tentativo)",
    "Cufe (tentativo)",
    "Fuente de los datos",
    "Archivo guardado",
    "Id mensaje",
]


def _abrir_o_crear(ruta: str) -> Workbook:
    if os.path.exists(ruta):
        return load_workbook(ruta)
    wb = Workbook()
    hoja = wb.active
    hoja.title = "Pendientes"
    for columna, encabezado in enumerate(ENCABEZADOS_PENDIENTES, start=1):
        hoja.cell(row=1, column=columna, value=encabezado)
    return wb


def agregar_pendiente(ruta: str, entrada: dict) -> None:
    """`entrada` trae las mismas claves que ENCABEZADOS_PENDIENTES, en snake_case:
    fecha_deteccion, hora_recepcion, remitente, asunto, categoria, motivo,
    factura, nit, proveedor, valor_sin_iva, cufe, fuente_datos, archivo_guardado,
    id_mensaje.
    """
    wb = _abrir_o_crear(ruta)
    hoja = wb["Pendientes"]

    fila = [
        entrada.get("fecha_deteccion"),
        entrada.get("hora_recepcion"),
        entrada.get("remitente"),
        entrada.get("asunto"),
        entrada.get("categoria"),
        entrada.get("motivo"),
        entrada.get("factura"),
        entrada.get("nit"),
        entrada.get("proveedor"),
        entrada.get("valor_sin_iva"),
        entrada.get("cufe"),
        entrada.get("fuente_datos"),
        entrada.get("archivo_guardado"),
        entrada.get("id_mensaje"),
    ]
    hoja.append(fila)

    excel_store.respaldar(ruta)
    wb.save(ruta)
