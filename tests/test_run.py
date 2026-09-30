import datetime as dt
import io
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook

from facturacion import config, excel_store, run
from facturacion.mail_client import Adjunto, CorreoRecibido

XML_VALIDO = b"""<?xml version="1.0" encoding="UTF-8"?>
<Invoice xmlns="urn:oasis:names:specification:ubl:schema:xsd:Invoice-2"
         xmlns:cac="urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2"
         xmlns:cbc="urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2">
    <cbc:ID>RUN001</cbc:ID>
    <cbc:UUID>cufe-run-001</cbc:UUID>
    <cbc:IssueDate>2026-09-08</cbc:IssueDate>
    <cac:AccountingSupplierParty>
        <cac:Party>
            <cac:PartyTaxScheme>
                <cbc:RegistrationName>Proveedor Run SAS</cbc:RegistrationName>
                <cbc:CompanyID>900333444</cbc:CompanyID>
            </cac:PartyTaxScheme>
        </cac:Party>
    </cac:AccountingSupplierParty>
    <cac:LegalMonetaryTotal>
        <cbc:TaxExclusiveAmount>80000.00</cbc:TaxExclusiveAmount>
        <cbc:PayableAmount>95200.00</cbc:PayableAmount>
    </cac:LegalMonetaryTotal>
</Invoice>
"""


def _zip_con_xml(xml_bytes: bytes) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as zf:
        zf.writestr("ad_run.xml", xml_bytes)
    return buffer.getvalue()


def _correo(asunto="Factura", adjuntos=None, id_mensaje="<run@test>", fecha_hora_utc=None):
    return CorreoRecibido(
        id_mensaje=id_mensaje,
        id_imap=b"1",
        asunto=asunto,
        remitente="proveedor@ejemplo.com",
        # 20:00 UTC == 15:00 hora Colombia (UTC-5). Se puede pisar con un
        # valor reciente cuando la prueba pasa por el filtro de ventana de
        # tiempo real de ejecutar_corrida (ver test_xml_corrupto_no_tumba...).
        fecha_hora_utc=fecha_hora_utc or dt.datetime(2026, 9, 8, 20, 0, 0),
        cuerpo_texto="",
        adjuntos=adjuntos or [],
    )


def _preparar_rutas(tmp_path, monkeypatch):
    carpeta_drive = tmp_path / "Facturas Houston"
    carpeta_drive.mkdir()
    ruta_excel = tmp_path / "Consolidado 2026.xlsx"
    ruta_cola = tmp_path / "Pendientes_de_revision.xlsx"

    monkeypatch.setattr(config, "CARPETA_DRIVE_LOCAL", str(carpeta_drive))
    monkeypatch.setattr(config, "RUTA_EXCEL_CONTROL", str(ruta_excel))
    monkeypatch.setattr(config, "RUTA_COLA_REVISION", str(ruta_cola))
    return carpeta_drive, ruta_excel, ruta_cola


def test_factura_clara_se_registra_y_guarda_archivo(tmp_path, monkeypatch):
    carpeta_drive, ruta_excel, _ = _preparar_rutas(tmp_path, monkeypatch)

    correo = _correo(adjuntos=[Adjunto("factura.zip", _zip_con_xml(XML_VALIDO))])
    wb = Workbook()
    resultado_texto = run._procesar_correo(correo, wb)

    assert "REGISTRADA" in resultado_texto
    hoja = wb["08-09-2026"]
    assert hoja.cell(row=2, column=config.COL_FACTURA).value == "RUN001"
    assert hoja.cell(row=2, column=config.COL_HORA_RECEPCION).value == dt.time(15, 0, 0)

    archivos_guardados = list((carpeta_drive / "08-09-2026").glob("*.zip"))
    assert len(archivos_guardados) == 1


def test_factura_duplicada_no_se_registra_dos_veces(tmp_path, monkeypatch):
    carpeta_drive, ruta_excel, _ = _preparar_rutas(tmp_path, monkeypatch)

    correo = _correo(adjuntos=[Adjunto("factura.zip", _zip_con_xml(XML_VALIDO))])
    wb = Workbook()
    run._procesar_correo(correo, wb)

    correo_reenviado = _correo(adjuntos=[Adjunto("factura.zip", _zip_con_xml(XML_VALIDO))], id_mensaje="<otro@test>")
    resultado_texto = run._procesar_correo(correo_reenviado, wb)

    assert "DUPLICADO" in resultado_texto
    hoja = wb["08-09-2026"]
    assert hoja.max_row == 2  # sigue habiendo una sola fila de datos


def test_caso_ambiguo_va_a_cola_de_revision(tmp_path, monkeypatch):
    carpeta_drive, ruta_excel, ruta_cola = _preparar_rutas(tmp_path, monkeypatch)

    correo = _correo(asunto="900111222;Proveedor SAS;NC-09;91;Proveedor SAS")
    wb = Workbook()
    resultado_texto = run._procesar_correo(correo, wb)

    assert "PENDIENTE_REVISION" in resultado_texto
    assert ruta_cola.exists()
    from openpyxl import load_workbook

    hoja_cola = load_workbook(ruta_cola)["Pendientes"]
    assert hoja_cola.cell(row=2, column=7).value == "NC-09"  # columna "Factura (tentativa)"


def test_xml_corrupto_no_tumba_la_corrida_y_guarda_el_zip_original(tmp_path, monkeypatch):
    """Caso real: un XML malformado dentro del zip no debe detener el resto de la corrida,
    y el zip original debe quedar disponible en la cola de revisión (no perderse)."""
    carpeta_drive, ruta_excel, ruta_cola = _preparar_rutas(tmp_path, monkeypatch)
    monkeypatch.setattr(run, "RUTA_ESTADO", str(tmp_path / "estado" / "procesados.json"))
    monkeypatch.setattr(run, "RUTA_LOG", str(tmp_path / "estado" / "log_corridas.txt"))

    hace_una_hora = dt.datetime.now(dt.timezone.utc).replace(tzinfo=None) - dt.timedelta(hours=1)
    zip_con_xml_corrupto = _zip_con_xml(b"esto no es un xml valido <<<")
    correo_malo = _correo(
        adjuntos=[Adjunto("factura.zip", zip_con_xml_corrupto)], id_mensaje="<malo@test>", fecha_hora_utc=hace_una_hora
    )
    correo_bueno = _correo(
        adjuntos=[Adjunto("factura.zip", _zip_con_xml(XML_VALIDO))], id_mensaje="<bueno@test>", fecha_hora_utc=hace_una_hora
    )

    class ConexionFalsa:
        def logout(self):
            pass

    monkeypatch.setattr(run.mail_client, "conectar", lambda: ConexionFalsa())
    monkeypatch.setattr(run.mail_client, "obtener_ids_en_ventana", lambda *a, **k: [b"1", b"2"])
    correos_por_id = {b"1": correo_malo, b"2": correo_bueno}
    monkeypatch.setattr(run.mail_client, "obtener_mensaje_completo", lambda conexion, id_imap: correos_por_id[id_imap])

    run.ejecutar_corrida()  # no debe lanzar excepción

    from openpyxl import load_workbook

    # La factura buena sí se registró a pesar del error en la otra.
    fecha_colombia = (hace_una_hora + run.OFFSET_COLOMBIA).date()
    wb_final = load_workbook(ruta_excel)
    nombre_pestana = fecha_colombia.strftime("%d-%m-%Y")
    assert wb_final[nombre_pestana].cell(row=2, column=config.COL_FACTURA).value == "RUN001"

    # El correo con XML corrupto quedó en la cola de revisión, con el zip original guardado.
    hoja_cola = load_workbook(ruta_cola)["Pendientes"]
    fila_error = [hoja_cola.cell(row=2, column=c).value for c in range(1, 15)]
    assert "Error inesperado" in fila_error[5]  # columna "Motivo de la duda"
    ruta_guardada = fila_error[12]  # columna "Archivo guardado"
    assert ruta_guardada and Path(ruta_guardada).exists()
